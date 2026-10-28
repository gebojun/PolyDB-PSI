"""Save an accurate Chinese caption for the finalized A-F figure."""

from pathlib import Path
import json
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


sys.stdout.reconfigure(encoding="utf-8")
OUTPUT = Path(__file__).resolve().parent.parent / "outputs/14_score_error_validation_cross_property"
report = json.loads((OUTPUT / "six_panel_figure_report.json").read_text(encoding="utf-8"))
if report["external_report_count"] != 7874 or report["panel_c_property_count"] != 31:
    raise ValueError("Caption and figure populations differ")

title = "图 3. PSI 的内部验证、结构距离分层分析及跨数据库外部验证。"
panels = [
    "（A）玻璃化转变温度（Tg）和介电常数（Dk）的严格折外预测按 PSI 从低到高划分为十个等频组（D1–D10），柱高为标准化绝对误差中位数。",
    "（B）按 PSI 从高到低优先保留预测（top，实线）与随机排序保留预测（rand，虚线）时，平均标准化绝对误差随保留比例的变化。",
    "（C）所示 31 项性质的标准化折外绝对误差与 PSI 的 Spearman 秩相关性，以 −ρ(e, PSI) 表示，颜色区分性质大类；每类优先展示折外样本数最多的至多五项性质，横坐标采用补充信息中的映射标签。",
    "（D）Tg 的结构距离四分位组（Q1–Q4，由近至远）内，不同 PSI 组的平均标准化绝对误差；High、Medium 和 Low PSI 分别对应全体 Tg 预测的 PSI 排名最高 20%、中间 60% 和最低 20%。",
    "（E）PolyDB 模型预测值与 PolyMetriX 逐条实验 Tg 报告的比较。颜色保留作者的 Gold、Yellow、Black 和 Red 标签，红色星形表示 Red 报告，斜虚线表示 y=x。",
    "（F）按全部 7,367 个唯一 pSMILES 的 PSI 四分位数划分的四个排名区间中，预测与实际实验报告的平均绝对误差；柱上数字分别为 MAE 和报告数 n，水平虚线为全部报告的平均误差（20.09 K）。",
]
notes = (
    "标准化绝对误差为绝对预测误差除以相应性质实验观测值的标准差。"
    "A 的误差棒为 500 次 bootstrap 的 95% 区间，D 为均值标准误（SEM），"
    "F 为固定模型及分组下 2,000 次按结构整组 bootstrap 的 95% 区间。"
    "E/F 均保留全部 7,874 条报告，包括来自四个结构的 30 条 Red 报告；"
    "该完整评估集包含与训练集共有的化学图结构。"
)
expanded_caption = title + "".join(panels) + notes
caption = (
    "图 3. PSI 的内部与跨数据库验证。"
    "（A）Tg 与 Dk 按 PSI 十等分后的标准化绝对误差中位数；"
    "（B）PSI 优先保留与随机保留的误差–覆盖率关系；"
    "（C）各性质预测误差与 PSI 的 Spearman 相关性（以 −ρ 表示）；"
    "（D）结构距离四分位组内不同 PSI 组的平均标准化误差；"
    "（E）PolyDB 模型预测与 PolyMetriX 实验 Tg 的比较，颜色为作者标签；"
    "（F）PSI 四分位组的平均外部误差。"
    "A/F 的误差棒为 95% bootstrap 区间，D 为 SEM。"
    "E/F 包含全部 7,874 条报告（Red=30）。"
)
method_note = (
    "D 沿用全体 Tg 的固定 PSI 20/60/20 排名分组，不是在各距离区间内重新划分 20/60/20。"
    "F 的四个区间互不重叠，四分位分界值为 "
    + "、".join(f"{x:.6f}" for x in report["external_psi_cutoffs"])
    + "，每个结构的全部报告继承同一 PSI 与分组。"
)
(OUTPUT / "图3_PSI验证六面板图注.md").write_text(
    caption + "\n",
    encoding="utf-8",
)
(OUTPUT / "图3_PSI验证图注_详细方法说明.md").write_text(
    expanded_caption + "\n\n绘图方法核对说明：" + method_note + "\n",
    encoding="utf-8",
)
document = Document()
section = document.sections[0]
section.left_margin = section.right_margin = Inches(.7)
normal = document.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.25
document.add_paragraph(caption)
document.add_picture(str(OUTPUT / "psi_validation_abcdef.png"), width=Inches(6.8))
target = OUTPUT / "图3_PSI验证六面板图注.docx"
document.save(target)
print(target)
print(caption)
